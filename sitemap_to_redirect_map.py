#!/usr/bin/env python3
"""
sitemap_to_redirect_map.py

Automates the first half of a site-migration redirect map:
  - reads a site's sitemap.xml (or sitemap index)
  - visits every URL and pulls a clean page title
  - writes an .xlsx in the "RW Task / 3D Task" layout, with
    Column A = Page title, Column B = Old Site URL, Column C = New Site URL (blank)

Usage:
    python sitemap_to_redirect_map.py https://www.oldsite.com/ -o redirect_map.xlsx
    python sitemap_to_redirect_map.py https://www.oldsite.com/sitemap.xml -o redirect_map.xlsx

Requires: requests, beautifulsoup4, lxml, openpyxl
    pip install requests beautifulsoup4 lxml openpyxl
"""

import argparse
import re
import sys
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; redirect-map-bot/1.0)"}
TIMEOUT = 15


def discover_wp_rest_pages(site_url: str, max_pages: int = 1000) -> list:
    """
    Query the WordPress REST API (/wp-json/wp/v2/pages) to enumerate EVERY
    published page, including orphaned ones with no incoming links anywhere
    on the site — something neither a sitemap.xml nor link-crawling can
    guarantee to catch. Works on any standard WordPress install where the
    REST API hasn't been disabled. Returns [] quietly if it's unavailable,
    blocked, or the site isn't WordPress.
    """
    parsed = urlparse(site_url)
    api_base = f"{parsed.scheme}://{parsed.netloc}/wp-json/wp/v2/pages"

    urls = []
    page = 1
    while True:
        try:
            r = requests.get(
                api_base,
                params={"per_page": 100, "page": page, "_fields": "link"},
                headers=HEADERS,
                timeout=TIMEOUT,
            )
        except requests.RequestException:
            return urls

        if r.status_code != 200:
            return urls  # REST API disabled, blocked, or not WordPress

        try:
            data = r.json()
        except ValueError:
            return urls

        if not isinstance(data, list) or not data:
            break

        urls.extend(item["link"] for item in data if item.get("link"))

        total_pages = r.headers.get("X-WP-TotalPages")
        if (total_pages and page >= int(total_pages)) or len(data) < 100 or len(urls) >= max_pages:
            break
        page += 1

    return urls


def find_sitemap_url(base_url: str) -> str:
    """If given a bare domain, locate its sitemap via robots.txt or common paths."""
    if base_url.rstrip("/").endswith(".xml"):
        return base_url

    parsed = urlparse(base_url)
    root = f"{parsed.scheme}://{parsed.netloc}/"

    # 1. Check robots.txt
    try:
        r = requests.get(urljoin(root, "robots.txt"), headers=HEADERS, timeout=TIMEOUT)
        if r.ok:
            m = re.search(r"(?im)^sitemap:\s*(\S+)", r.text)
            if m:
                return m.group(1).strip()
    except requests.RequestException:
        pass

    # 2. Common fallback paths
    for path in ("sitemap.xml", "sitemap_index.xml", "wp-sitemap.xml"):
        candidate = urljoin(root, path)
        try:
            r = requests.get(candidate, headers=HEADERS, timeout=TIMEOUT)
            if r.ok and "<url" in r.text.lower():
                return candidate
        except requests.RequestException:
            continue

    raise RuntimeError(f"Couldn't locate a sitemap for {base_url}. Pass the sitemap.xml URL directly.")


def collect_urls(sitemap_url: str, seen=None) -> list:
    """Recursively expand sitemap indexes into a flat list of page URLs."""
    if seen is None:
        seen = set()
    if sitemap_url in seen:
        return []
    seen.add(sitemap_url)

    r = requests.get(sitemap_url, headers=HEADERS, timeout=TIMEOUT)
    r.raise_for_status()
    soup = BeautifulSoup(r.content, "xml")

    # Sitemap index: contains <sitemap><loc> entries pointing at other sitemaps
    sub_sitemaps = [loc.get_text(strip=True) for loc in soup.select("sitemap > loc")]
    if sub_sitemaps:
        urls = []
        for sm in sub_sitemaps:
            urls.extend(collect_urls(sm, seen))
        return urls

    # Regular sitemap: <url><loc> entries are pages
    return [loc.get_text(strip=True) for loc in soup.select("url > loc")]


# File extensions that are never real "pages" and shouldn't be followed while crawling
_ASSET_EXT_RE = re.compile(
    r"\.(jpg|jpeg|png|gif|svg|webp|ico|bmp|pdf|zip|rar|css|js|mp4|mov|avi|mp3|wav|woff2?|ttf|eot|xml|json)$",
    re.IGNORECASE,
)


def _canonicalize(url: str) -> str:
    """
    Collapse http vs https, trailing-slash differences, and netloc case into
    one key, so the same page reached by two different URLs is only crawled
    and recorded once.
    """
    p = urlparse(url)
    netloc = p.netloc.lower()
    path = p.path.rstrip("/")
    query = f"?{p.query}" if p.query else ""
    return f"{netloc}{path}{query}"


def _extract_canonical(soup: BeautifulSoup, response_url: str) -> str:
    """
    Return the URL a page declares as its own canonical address via
    <link rel="canonical">, resolved against response_url. Falls back to
    response_url itself if there's no canonical tag.
    """
    tag = soup.find("link", rel="canonical")
    href = tag.get("href", "").strip() if tag else ""
    return urljoin(response_url, href) if href else response_url


def crawl_site(start_url: str, max_pages: int = 300) -> list:
    """
    Fallback for sites with no discoverable sitemap: starting from the homepage,
    follow same-domain <a href> links breadth-first and collect page URLs.
    Common on custom-themed WordPress builds (e.g. Proweaver sites) that don't
    run an SEO plugin's sitemap feature.

    Hierarchical WordPress pages are frequently reachable at more than one
    path -- e.g. a nav link to the flat "/child-page/" for content that is
    actually a child of "/parent-page/", where "/parent-page/child-page/"
    also resolves. WordPress itself always knows the "real" address and
    stamps it onto the page as <link rel="canonical">, so every fetched page
    is resolved against its own canonical URL before being recorded. That
    both (a) captures the correct full path instead of a flat shortcut link,
    and (b) collapses duplicate pages reachable via two different hrefs down
    to a single row instead of recording each path separately.
    """
    root_netloc = urlparse(start_url).netloc.lower()
    to_visit = [start_url]
    queued = {_canonicalize(start_url)}
    seen = set()
    seen_resolved = set()
    found = []

    while to_visit and len(found) < max_pages:
        url = to_visit.pop(0)
        key = _canonicalize(url)
        if key in seen:
            continue
        seen.add(key)

        try:
            r = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
            r.raise_for_status()
        except requests.RequestException:
            continue

        soup = BeautifulSoup(r.text, "html.parser")

        resolved_url = _extract_canonical(soup, r.url)
        if urlparse(resolved_url).netloc.lower() != root_netloc:
            resolved_url = r.url  # ignore a canonical tag pointing off-site

        resolved_key = _canonicalize(resolved_url)
        if resolved_key not in seen_resolved:
            seen_resolved.add(resolved_key)
            found.append(resolved_url)
        seen.add(resolved_key)  # don't re-crawl this same page under its canonical URL either

        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if not href or href.startswith(("#", "mailto:", "tel:", "javascript:")):
                continue

            full = urljoin(r.url, href).split("#")[0]
            parsed_full = urlparse(full)

            if parsed_full.scheme not in ("http", "https"):
                continue
            if parsed_full.netloc.lower() != root_netloc:
                continue
            if _ASSET_EXT_RE.search(parsed_full.path):
                continue

            key_full = _canonicalize(full)
            if key_full not in seen and key_full not in queued:
                queued.add(key_full)
                to_visit.append(full)

    return found


# Pages that should never appear in the redirect map (matched against the URL path)
EXCLUDED_PATH_PATTERNS = (
    r"^sitemap/?$",
    r"^sitemap\.xml$",
    r"^sitemap[-_].*$",
    r"^category/.*$",
    r"^tag/.*$",
    r"^author/.*$",
)

# File extensions excluded from the redirect map wholesale
EXCLUDED_EXTENSIONS = (".php", ".pdf")

# Path segments that mark a URL as excluded wholesale, wherever they appear
EXCLUDED_SEGMENTS = ("feed",)


def is_excluded_url(url: str) -> bool:
    parsed = urlparse(url)
    path = parsed.path.strip("/").lower()

    if any(re.match(pattern, path) for pattern in EXCLUDED_PATH_PATTERNS):
        return True
    if path.endswith(EXCLUDED_EXTENSIONS):
        return True

    segments = path.split("/") if path else []
    if any(seg in EXCLUDED_SEGMENTS for seg in segments):
        return True

    # Exclude individual blog posts (anything nested under a "blog" segment),
    # but keep the blog index page itself (e.g. .../resources/blog).
    if "blog" in segments and segments.index("blog") != len(segments) - 1:
        return True

    if "feed" in parsed.query.lower():
        return True

    return False


def is_homepage_url(url: str, site_root: str) -> bool:
    path = urlparse(url).path.strip("/")
    return path == "" or url.rstrip("/") == site_root.rstrip("/")


# Matches a Yoast-style pipe separator: " | "
# (handles a stray non-breaking space too, which some SEO plugins insert)
_SEP_RE = r"[\s\xa0]*\|[\s\xa0]*"
_SEP_SPLIT_RE = re.compile(_SEP_RE)


def clean_title(raw_title: str) -> str:
    """
    Strip everything from the first pipe ("|") separator onward,
    e.g. "About Us | Home Care in Ohio | My Love Home Care LLC" -> "About Us".
    Titles that use a dash/hyphen separator are left untouched.
    """
    if not raw_title:
        return ""
    title = raw_title.strip()
    return _SEP_SPLIT_RE.split(title, maxsplit=1)[0].strip()


def _get_meta_content(soup: BeautifulSoup, prop_name: str) -> str:
    tag = soup.find("meta", attrs={"property": prop_name})
    content = tag.get("content", "").strip() if tag else ""
    return content


def extract_fallback_heading(soup: BeautifulSoup) -> str:
    """
    Return the first on-page heading (h1, then h2, then h3) that isn't just
    the site's own branding -- e.g. a logo wrapped in <h1>{Site Name}</h1>.
    Used as a fallback when a site's <title> tag is broken (some
    WordPress/theme setups emit the exact same <title> on every single page,
    which a plain title-tag scrape can't tell apart).
    """
    site_name = _get_meta_content(soup, "og:site_name").lower()
    for level in ("h1", "h2", "h3"):
        for tag in soup.find_all(level):
            text = tag.get_text(strip=True)
            if text and text.lower() != site_name:
                return text
    return ""


def fetch_page_data(url: str) -> tuple:
    """
    Fetch the page and return (raw_title, is_blog_post, fallback_heading).

    is_blog_post is detected via the <meta property="article:published_time">
    tag, which WordPress/Yoast adds to actual blog posts but never to regular
    pages (pages only ever get article:modified_time, if anything). This
    works even when posts live at a flat top-level URL with no /blog/ in the
    path, where a URL-pattern check alone can't tell them apart.

    fallback_heading is the page's on-page heading text (see
    extract_fallback_heading), used by resolve_duplicate_titles() when a
    site's <title> tag turns out to be identical across many pages.
    """
    try:
        r = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        is_post = soup.find("meta", attrs={"property": "article:published_time"}) is not None
        heading = extract_fallback_heading(soup)
        if soup.title and soup.title.string:
            return soup.title.string.strip(), is_post, heading
        return "", is_post, heading
    except requests.RequestException:
        pass
    # Fallback: derive a readable title from the URL slug (can't check post status)
    path = urlparse(url).path.strip("/")
    slug = path.split("/")[-1] if path else "Home"
    return (slug.replace("-", " ").replace("_", " ").title() or "Home"), False, ""


def resolve_duplicate_titles(entries: list) -> None:
    """
    entries: list of dicts with keys "title", "url", "heading" (mutated in
    place). If the same title shows up across more than one URL, it's
    almost certainly a site-wide <title> tag bug rather than genuinely
    identical pages -- so swap in that page's on-page heading instead,
    when one is available and actually differs from the broken title.
    """
    from collections import Counter
    counts = Counter(e["title"] for e in entries if e["title"] != "Home")
    for e in entries:
        if e["title"] == "Home" or counts[e["title"]] <= 1:
            continue
        heading = e.get("heading") or ""
        cleaned = clean_title(heading)
        if cleaned and cleaned.lower() != e["title"].lower():
            print(f"  Duplicate title detected for {e['url']} "
                  f"-> using on-page heading \"{cleaned}\" instead", file=sys.stderr)
            e["title"] = cleaned


def build_workbook(site_url: str, rows: list, out_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Redirect Map"

    blue = PatternFill("solid", fgColor="4285F4")
    gray = PatternFill("solid", fgColor="808080")
    white_bold = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center")

    # Row 1: site URL (hyperlinked) | RW Task | 3D Task
    ws["A1"] = site_url
    ws["A1"].hyperlink = site_url
    ws["A1"].font = Font(color="0563C1", underline="single")
    ws["B1"] = "RW Task"
    ws["C1"] = "3D Task"
    for cell in ("B1", "C1"):
        ws[cell].fill = blue
        ws[cell].font = white_bold
        ws[cell].alignment = center

    # Row 2: "N Pages" | Old Site | New Site
    # A2 counts non-empty title cells from row 3 down (data rows only — starting
    # here avoids counting A2's own formula cell, which would be circular),
    # so it stays accurate if rows are ever added/removed by hand later.
    ws["A2"] = '=COUNTA(A3:A1048576)&" Pages"'
    ws["B2"] = "Old Site"
    ws["C2"] = "New Site"
    for cell in ("A2", "B2", "C2"):
        ws[cell].fill = gray
        ws[cell].font = white_bold
        ws[cell].alignment = center

    # Data rows
    for i, (title, url) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=title)
        c = ws.cell(row=i, column=2, value=url)
        c.hyperlink = url
        c.font = Font(color="0563C1", underline="single")
        # Column C (new site) intentionally left blank

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60

    wb.save(out_path)


def main():
    try:
        import lxml  # noqa: F401
    except ImportError:
        sys.exit(
            "Missing dependency: 'lxml' is required to parse sitemap.xml files.\n"
            "Install it with:  pip install lxml"
        )

    ap = argparse.ArgumentParser(description="Build a redirect-map spreadsheet from a site's sitemap.")
    ap.add_argument("site", help="Old site's base URL or sitemap.xml URL")
    ap.add_argument("-o", "--output", default="redirect_map.xlsx", help="Output .xlsx path")
    ap.add_argument("--max-pages", type=int, default=300,
                     help="Max pages to visit when crawling a site with no sitemap (default: 300)")
    ap.add_argument("--no-crawl", action="store_true",
                     help="Don't fall back to crawling the site if no sitemap is found")
    ap.add_argument("--no-rest-api", action="store_true",
                     help="Don't supplement results with the WordPress REST API "
                          "(/wp-json/wp/v2/pages), which catches orphaned pages "
                          "that have no incoming links")
    ap.add_argument("--extra-urls", nargs="*", default=[],
                     help="Extra page URLs to always include, e.g. orphaned pages with no "
                          "incoming links that a crawl alone can't discover: "
                          "--extra-urls https://site.com/page-a https://site.com/page-b")
    args = ap.parse_args()

    print(f"Locating sitemap for {args.site} ...", file=sys.stderr)
    try:
        sitemap_url = find_sitemap_url(args.site)
        print(f"Using sitemap: {sitemap_url}", file=sys.stderr)
        all_urls = collect_urls(sitemap_url)
    except RuntimeError as e:
        if args.no_crawl:
            raise
        print(f"{e}", file=sys.stderr)
        print("Falling back to crawling the site for pages instead...", file=sys.stderr)
        all_urls = crawl_site(args.site, max_pages=args.max_pages)
        print(f"Crawled {len(all_urls)} page(s).", file=sys.stderr)

    if not args.no_rest_api:
        rest_urls = discover_wp_rest_pages(args.site)
        if rest_urls:
            existing_keys = {_canonicalize(u) for u in all_urls}
            added = [u for u in rest_urls if _canonicalize(u) not in existing_keys]
            all_urls = all_urls + added
            if added:
                print(f"WordPress REST API found {len(added)} additional page(s) "
                      f"(likely orphaned, no incoming links).", file=sys.stderr)

    if args.extra_urls:
        existing_keys = {_canonicalize(u) for u in all_urls}
        added = [u for u in args.extra_urls if _canonicalize(u) not in existing_keys]
        all_urls = all_urls + added
        if added:
            print(f"Added {len(added)} manually-specified extra URL(s).", file=sys.stderr)

    urls = [u for u in all_urls if not is_excluded_url(u)]
    skipped = len(all_urls) - len(urls)
    if skipped:
        print(f"Skipping {skipped} excluded URL(s) (.php/.pdf, feeds, blog posts, /sitemap).", file=sys.stderr)
    print(f"Found {len(urls)} URLs. Fetching titles...", file=sys.stderr)

    parsed = urlparse(args.site)
    site_root = f"{parsed.scheme}://{parsed.netloc}/"

    entries = []
    skipped_posts = 0
    for i, url in enumerate(urls, 1):
        if is_homepage_url(url, site_root):
            title, heading = "Home", None
        else:
            raw_title, is_post, heading = fetch_page_data(url)
            if is_post:
                skipped_posts += 1
                print(f"  [{i}/{len(urls)}] (blog post, skipped) -> {url}", file=sys.stderr)
                continue
            title = clean_title(raw_title)
        entries.append({"title": title, "url": url, "heading": heading})
        print(f"  [{i}/{len(urls)}] {title} -> {url}", file=sys.stderr)

    if skipped_posts:
        print(f"Skipped {skipped_posts} individual blog post(s).", file=sys.stderr)

    resolve_duplicate_titles(entries)
    rows = [(e["title"], e["url"]) for e in entries]

    build_workbook(site_root, rows, args.output)
    print(f"Done. Wrote {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()